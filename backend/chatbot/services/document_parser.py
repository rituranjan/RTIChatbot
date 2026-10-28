import csv
import json
from pathlib import Path

from .ollama_service import OllamaService


def extract_text(path) -> str:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path)
    if ext == ".xlsx":
        return _xlsx(path)
    if ext == ".csv":
        return _csv(path)
    if ext in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".json":
        return _json(path)
    if ext in (".png", ".jpg", ".jpeg", ".webp"):
        return _image(path)
    raise ValueError(f"Unsupported file type: {ext}")


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for sheet in wb.worksheets:
        out.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _csv(path: Path) -> str:
    with path.open(newline="", encoding="utf-8", errors="replace") as f:
        return "\n".join(" | ".join(row) for row in csv.reader(f))


def _json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    return json.dumps(data, indent=2, ensure_ascii=False)


def _image(path: Path) -> str:
    try:
        import pytesseract
        from PIL import Image

        text = pytesseract.image_to_string(Image.open(str(path)))
        if text.strip():
            return text
    except Exception:
        pass

    try:
        desc = OllamaService().describe_image(str(path))
        if desc:
            return desc
    except Exception:
        pass

    raise ValueError(
        "No OCR engine or vision model is available to read this image."
    )